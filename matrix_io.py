"""Read and write Caterpillar Air matrix workbooks."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_matrix import (
    BOLD_FONT,
    COLUMN_HEADER_ROW,
    COST_NAME_ROW,
    DATA_START_ROW,
    HEADER_FILL,
    HEADER_FONT,
    PRICE_NUMBER_FORMAT,
    RATE_BY_ROW,
    SHIPMENT_COLUMNS,
    SUBHEADER_FILL,
    SUBHEADER_FONT,
    THIN_BORDER,
)

EXCEL_SUFFIXES = {".xlsx", ".xls", ".xlsm"}
UPDATED_FILL_COLOR = "C6EFCE"
UPDATED_ROW_FILL_COLOR = "FFEB9C"
DEFAULT_MATRIX_SHEET = "Matrix"
RATE_CARD_SHEET_NAME = "Rate card"
RATE_CARD_HEADER_MARKERS = ("Lane #", "KEY")
RATE_CARD_COST_SKIP_PREFIXES = (
    "Applies if",
    "Rate by:",
    "Validity period",
    "Conditional rules",
)


@dataclass(frozen=True)
class MatrixColumn:
    index: int
    header: str
    cost_name: str | None = None


@dataclass
class MatrixLane:
    values: list[object]
    changed_columns: set[int] = field(default_factory=set)
    is_updated_row: bool = False


@dataclass
class MatrixWorkbookData:
    sheet_name: str
    columns: list[MatrixColumn]
    header_rows: list[list[object]]
    lanes: list[MatrixLane]
    extra_sheets: dict[str, pd.DataFrame] = field(default_factory=dict)
    header_row: int = COLUMN_HEADER_ROW
    data_start_row: int = DATA_START_ROW


def _updated_fill() -> PatternFill:
    return PatternFill(
        fill_type="solid",
        start_color=UPDATED_FILL_COLOR,
        end_color=UPDATED_FILL_COLOR,
    )


def _updated_row_fill() -> PatternFill:
    return PatternFill(
        fill_type="solid",
        start_color=UPDATED_ROW_FILL_COLOR,
        end_color=UPDATED_ROW_FILL_COLOR,
    )


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value).strip()


def _is_empty_value(value: object) -> bool:
    return _cell_text(value) == ""


def parse_date_dd_mm_yyyy(value: object) -> date | None:
    if _is_empty_value(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return None
    return parsed.date()


def format_date_dd_mm_yyyy(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def list_excel_files(directory: Path) -> list[Path]:
    return [
        path
        for path in sorted(directory.iterdir())
        if path.is_file()
        and path.suffix.lower() in EXCEL_SUFFIXES
        and not path.name.startswith("~$")
    ]


def prompt_file_selection(
    files: list[Path],
    title: str,
    directory: Path | None = None,
    *,
    always_prompt: bool = False,
) -> Path:
    if not files:
        location = directory or "the selected folder"
        raise FileNotFoundError(f"No Excel files found in: {location}")

    if len(files) == 1 and not always_prompt:
        print(f"\nUsing only file: {files[0].name}")
        return files[0]

    location = directory or files[0].parent
    print(f"\n{title}")
    print(f"  Folder: {location}")
    for index, path in enumerate(files, start=1):
        print(f"  {index}. {path.name}")

    while True:
        raw = input("Enter file number: ").strip()
        try:
            choice = int(raw) - 1
            if 0 <= choice < len(files):
                return files[choice]
        except ValueError:
            pass
        print(f"Please enter a number between 1 and {len(files)}.")


def _resolve_matrix_sheet_name(workbook) -> str:
    if DEFAULT_MATRIX_SHEET in workbook.sheetnames:
        return DEFAULT_MATRIX_SHEET
    return workbook.sheetnames[0]


def resolve_rate_card_sheet_name(workbook) -> str:
    if RATE_CARD_SHEET_NAME in workbook.sheetnames:
        return RATE_CARD_SHEET_NAME
    raise ValueError(
        f"Sheet '{RATE_CARD_SHEET_NAME}' not found. "
        f"Available: {', '.join(workbook.sheetnames)}"
    )


def _read_header_rows(worksheet, column_count: int, last_row: int = COLUMN_HEADER_ROW) -> list[list[object]]:
    header_rows: list[list[object]] = []
    for row_index in range(1, last_row + 1):
        row_values: list[object] = []
        for column_index in range(1, column_count + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            row_values.append(None if value == "" else value)
        header_rows.append(row_values)
    return header_rows


def _find_rate_card_header_row(worksheet) -> int:
    for row_index in range(1, min(worksheet.max_row, 50) + 1):
        row_values = [
            _cell_text(worksheet.cell(row=row_index, column=column_index).value)
            for column_index in range(1, min(worksheet.max_column, 40) + 1)
        ]
        if RATE_CARD_HEADER_MARKERS[0] in row_values and RATE_CARD_HEADER_MARKERS[1] in row_values:
            return row_index
    raise ValueError(
        f"Could not find '{RATE_CARD_HEADER_MARKERS[0]}' and "
        f"'{RATE_CARD_HEADER_MARKERS[1]}' header row in '{RATE_CARD_SHEET_NAME}'."
    )


def _rate_card_cost_name(worksheet, column_index: int, header_row: int) -> str | None:
    for row_index in range(1, header_row):
        value = worksheet.cell(row=row_index, column=column_index).value
        if _is_empty_value(value):
            continue
        text = str(value).strip().splitlines()[0].strip()
        if not text or text in {"MIN", "<= 10000"}:
            continue
        if any(text.startswith(prefix) for prefix in RATE_CARD_COST_SKIP_PREFIXES):
            continue
        return text
    return None


def _build_rate_card_columns(worksheet, header_row: int, column_count: int) -> list[MatrixColumn]:
    columns: list[MatrixColumn] = []
    current_cost_name: str | None = None

    for index in range(column_count):
        column_number = index + 1
        header_text = _cell_text(worksheet.cell(row=header_row, column=column_number).value)
        if header_text == "Currency":
            current_cost_name = _rate_card_cost_name(worksheet, column_number, header_row)
            columns.append(
                MatrixColumn(index=index, header=header_text, cost_name=current_cost_name)
            )
        elif current_cost_name is not None and header_text:
            columns.append(
                MatrixColumn(index=index, header=header_text, cost_name=current_cost_name)
            )
        else:
            current_cost_name = None
            columns.append(MatrixColumn(index=index, header=header_text, cost_name=None))

    return columns


def _build_columns(header_rows: list[list[object]]) -> list[MatrixColumn]:
    cost_name_row = header_rows[COST_NAME_ROW - 1]
    header_row = header_rows[COLUMN_HEADER_ROW - 1]
    columns: list[MatrixColumn] = []
    for index, header in enumerate(header_row):
        header_text = _cell_text(header)
        cost_name = _cell_text(cost_name_row[index]) or None
        columns.append(MatrixColumn(index=index, header=header_text, cost_name=cost_name))
    return columns


def read_matrix_workbook(
    file_path: Path,
    *,
    sheet_name: str | None = None,
    rate_card_only: bool = False,
) -> MatrixWorkbookData:
    workbook = load_workbook(file_path, data_only=True)
    if rate_card_only:
        resolved_sheet = resolve_rate_card_sheet_name(workbook)
    else:
        resolved_sheet = sheet_name or _resolve_matrix_sheet_name(workbook)
    worksheet = workbook[resolved_sheet]

    column_count = worksheet.max_column or 0
    if rate_card_only:
        header_row = _find_rate_card_header_row(worksheet)
        data_start_row = header_row + 1
        header_rows = _read_header_rows(worksheet, column_count, last_row=header_row)
        columns = _build_rate_card_columns(worksheet, header_row, column_count)
    else:
        header_row = COLUMN_HEADER_ROW
        data_start_row = DATA_START_ROW
        header_rows = _read_header_rows(worksheet, column_count)
        columns = _build_columns(header_rows)

    lanes: list[MatrixLane] = []
    for row_index in range(data_start_row, worksheet.max_row + 1):
        row_values: list[object] = []
        has_data = False
        for column_index in range(1, column_count + 1):
            value = worksheet.cell(row=row_index, column=column_index).value
            if value == "":
                value = None
            if not _is_empty_value(value):
                has_data = True
            row_values.append(value)
        if has_data:
            lanes.append(MatrixLane(values=row_values))

    extra_sheets: dict[str, pd.DataFrame] = {}
    if not rate_card_only:
        for name in workbook.sheetnames:
            if name == resolved_sheet:
                continue
            extra_sheets[name] = pd.read_excel(file_path, sheet_name=name)

    return MatrixWorkbookData(
        sheet_name=resolved_sheet,
        columns=columns,
        header_rows=header_rows,
        lanes=lanes,
        extra_sheets=extra_sheets,
        header_row=header_row,
        data_start_row=data_start_row,
    )


def lane_value(lane: MatrixLane, columns: list[MatrixColumn], header: str) -> object:
    for column in columns:
        if column.header == header and column.cost_name is None:
            return lane.values[column.index]
    return None


def cell_text(value: object) -> str:
    return _cell_text(value)


def group_lanes_by_key(
    lanes: list[MatrixLane],
    columns: list[MatrixColumn],
) -> dict[str, list[MatrixLane]]:
    grouped: dict[str, list[MatrixLane]] = {}
    for lane in lanes:
        key = cell_text(lane_value(lane, columns, "KEY"))
        if not key:
            continue
        grouped.setdefault(key, []).append(lane)
    return grouped


def set_lane_value(
    lane: MatrixLane,
    columns: list[MatrixColumn],
    header: str,
    value: object,
) -> None:
    for column in columns:
        if column.header == header and column.cost_name is None:
            lane.values[column.index] = value
            return
    raise KeyError(f"Shipment column not found: {header}")


def find_cost_column_index(
    columns: list[MatrixColumn],
    cost_name: str,
    header: str,
) -> int | None:
    for column in columns:
        if column.cost_name == cost_name and column.header == header:
            return column.index
    return None


def normalize_cost_name(cost_name: str) -> str:
    text = cost_name.strip()
    if text.endswith(" Rate"):
        text = text[: -len(" Rate")]
    if "(" in text:
        text = text.split("(", 1)[0].strip()
    return text


def _expand_cost_name_aliases(cost_name: str) -> set[str]:
    normalized = normalize_cost_name(cost_name).lower()
    aliases = {normalized}
    if normalized in {"inspection fee", "origin cargo inspection"}:
        aliases.update({"inspection fee", "origin cargo inspection"})
    return aliases


def _header_match_priority(update_header: str) -> list[str]:
    if update_header == "Currency":
        return ["Currency"]
    if update_header == "Min":
        return ["MIN", "Min"]
    if update_header == "Max":
        return ["Max"]
    if update_header == "Flat":
        return ["Flat", "p/unit"]
    return [update_header]


def find_cost_column_index_fuzzy(
    columns: list[MatrixColumn],
    cost_name: str,
    header: str,
    *,
    used_indices: set[int] | None = None,
) -> int | None:
    used_indices = used_indices or set()
    exact = find_cost_column_index(columns, cost_name, header)
    if exact is not None and exact not in used_indices:
        return exact

    name_candidates = _expand_cost_name_aliases(cost_name)
    matching_columns = [
        column
        for column in columns
        if column.cost_name is not None
        and column.index not in used_indices
        and normalize_cost_name(column.cost_name).lower() in name_candidates
    ]

    for preferred_header in _header_match_priority(header):
        for column in matching_columns:
            if column.header == preferred_header:
                return column.index
    return None


def build_cost_column_map(
    target_columns: list[MatrixColumn],
    update_columns: list[MatrixColumn],
) -> dict[int, int]:
    mapping: dict[int, int] = {}
    used_target_indices: set[int] = set()

    for update_column in update_columns:
        if update_column.cost_name is None:
            continue
        target_index = find_cost_column_index_fuzzy(
            target_columns,
            update_column.cost_name,
            update_column.header,
            used_indices=used_target_indices,
        )
        if target_index is None:
            continue
        mapping[update_column.index] = target_index
        used_target_indices.add(target_index)

    return mapping


def find_shipment_column_index(columns: list[MatrixColumn], header: str) -> int | None:
    for column in columns:
        if column.cost_name is None and column.header == header:
            return column.index
    return None


def append_cost_columns(
    matrix: MatrixWorkbookData,
    cost_name: str,
    sub_headers: list[str],
) -> list[int]:
    added_indices: list[int] = []
    for sub_header in sub_headers:
        index = len(matrix.columns)
        matrix.columns.append(
            MatrixColumn(index=index, header=sub_header, cost_name=cost_name)
        )
        for row_index, header_row in enumerate(matrix.header_rows):
            if row_index == 0:
                header_row.append(cost_name)
            elif row_index == 1:
                header_row.append("")
            else:
                header_row.append(sub_header)
        for lane in matrix.lanes:
            lane.values.append(None)
        added_indices.append(index)
    return added_indices


def find_last_column_of_cost_block(
    columns: list[MatrixColumn],
    anchor_cost_name: str,
) -> int | None:
    anchor_key = normalize_cost_name(anchor_cost_name).lower()
    matching = [
        column.index
        for column in columns
        if column.cost_name is not None
        and normalize_cost_name(column.cost_name).lower() == anchor_key
    ]
    return max(matching) if matching else None


def find_cost_block_columns(
    columns: list[MatrixColumn],
    cost_name: str,
) -> list[MatrixColumn]:
    return [column for column in columns if column.cost_name == cost_name]


def _reindex_matrix_columns(matrix: MatrixWorkbookData) -> None:
    matrix.columns = [
        MatrixColumn(index=index, header=column.header, cost_name=column.cost_name)
        for index, column in enumerate(matrix.columns)
    ]


def find_insert_position_for_periodical_block(
    columns: list[MatrixColumn],
    after_anchor: str,
    display_base_name: str,
) -> int:
    anchor_end = find_last_column_of_cost_block(columns, after_anchor)
    if anchor_end is None:
        raise ValueError(f"Could not find cost block '{after_anchor}' to insert after.")

    insert_pos = anchor_end
    prefix = display_base_name.lower()
    for column in columns:
        if column.cost_name and column.cost_name.lower().startswith(prefix):
            insert_pos = max(insert_pos, column.index)
    return insert_pos + 1


def insert_cost_block_after(
    matrix: MatrixWorkbookData,
    after_anchor: str,
    cost_name: str,
    sub_headers: list[str],
    header_cells_by_row: list[list[object]],
    *,
    display_base_name: str | None = None,
) -> list[int]:
    existing = find_cost_block_columns(matrix.columns, cost_name)
    if existing:
        return [column.index for column in existing]

    if display_base_name is not None:
        insert_pos = find_insert_position_for_periodical_block(
            matrix.columns,
            after_anchor,
            display_base_name,
        )
    else:
        anchor_end = find_last_column_of_cost_block(matrix.columns, after_anchor)
        if anchor_end is None:
            raise ValueError(f"Could not find cost block '{after_anchor}' to insert after.")
        insert_pos = anchor_end + 1
    added_indices: list[int] = []

    for offset, sub_header in enumerate(sub_headers):
        matrix.columns.insert(
            insert_pos + offset,
            MatrixColumn(index=insert_pos + offset, header=sub_header, cost_name=cost_name),
        )
        for row_index, header_row in enumerate(matrix.header_rows):
            row_values = header_cells_by_row[row_index] if row_index < len(header_cells_by_row) else []
            header_row.insert(
                insert_pos + offset,
                row_values[offset] if offset < len(row_values) else sub_header,
            )
        for lane in matrix.lanes:
            lane.values.insert(insert_pos + offset, None)
        added_indices.append(insert_pos + offset)

    _reindex_matrix_columns(matrix)
    return added_indices


def _shipment_column_count(columns: list[MatrixColumn]) -> int:
    return sum(1 for column in columns if column.cost_name is None)


def _discover_cost_blocks_from_columns(columns: list[MatrixColumn]) -> list[tuple[str, list[MatrixColumn]]]:
    shipment_count = _shipment_column_count(columns)
    blocks: list[tuple[str, list[MatrixColumn]]] = []
    current_name: str | None = None
    current_columns: list[MatrixColumn] = []

    for column in columns[shipment_count:]:
        if column.cost_name != current_name:
            if current_name is not None:
                blocks.append((current_name, current_columns))
            current_name = column.cost_name
            current_columns = [column]
        else:
            current_columns.append(column)

    if current_name is not None:
        blocks.append((current_name, current_columns))
    return blocks


def _apply_worksheet_formatting(
    worksheet,
    columns: list[MatrixColumn],
    total_rows: int,
    *,
    changed_cells: set[tuple[int, int]],
) -> None:
    shipment_column_count = _shipment_column_count(columns)
    total_columns = len(columns)
    header_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    bold_shipment_headers = {
        column.header
        for column in SHIPMENT_COLUMNS
        if column.bold_header
    }

    for row_index in range(1, total_rows + 1):
        for column_index in range(1, total_columns + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.border = THIN_BORDER
            if (row_index, column_index) in changed_cells:
                cell.fill = _updated_fill()

            if row_index == COST_NAME_ROW and column_index > shipment_column_count:
                if (row_index, column_index) not in changed_cells:
                    cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = header_center
            elif row_index == RATE_BY_ROW:
                cell.alignment = center
            elif row_index == COLUMN_HEADER_ROW:
                if (row_index, column_index) not in changed_cells:
                    cell.fill = SUBHEADER_FILL
                header_font = SUBHEADER_FONT
                if column_index <= shipment_column_count:
                    column = columns[column_index - 1]
                    if column.header in bold_shipment_headers:
                        header_font = BOLD_FONT
                cell.font = header_font
                cell.alignment = center
            elif row_index >= DATA_START_ROW and column_index <= shipment_column_count:
                cell.alignment = left
            elif row_index >= DATA_START_ROW and column_index > shipment_column_count:
                cell.alignment = center
                cell.number_format = PRICE_NUMBER_FORMAT


def _write_matrix_sheet(
    worksheet,
    matrix: MatrixWorkbookData,
    *,
    changed_cells: set[tuple[int, int]] | None = None,
) -> None:
    total_columns = len(matrix.columns)
    for row_offset, header_row in enumerate(matrix.header_rows):
        for column_index, value in enumerate(header_row, start=1):
            cell = worksheet.cell(row=row_offset + 1, column=column_index, value=value)
            if value == "":
                cell.value = None

    for lane_offset, lane in enumerate(matrix.lanes):
        row_index = matrix.data_start_row + lane_offset
        for column_index, value in enumerate(lane.values, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if value == "":
                cell.value = None

    shipment_column_count = _shipment_column_count(matrix.columns)
    cost_blocks = _discover_cost_blocks_from_columns(matrix.columns)
    current_column = shipment_column_count + 1
    for _, block_columns in cost_blocks:
        block_width = len(block_columns)
        if block_width > 1:
            worksheet.merge_cells(
                start_row=COST_NAME_ROW,
                start_column=current_column,
                end_row=COST_NAME_ROW,
                end_column=current_column + block_width - 1,
            )
        current_column += block_width

    lane_number_index = find_shipment_column_index(matrix.columns, "Lane #")
    if lane_number_index is not None:
        for lane_offset, lane in enumerate(matrix.lanes):
            lane.values[lane_number_index] = lane_offset + 1

    total_rows = matrix.data_start_row + len(matrix.lanes) - 1
    _apply_worksheet_formatting(
        worksheet,
        matrix.columns,
        total_rows,
        changed_cells=changed_cells or set(),
    )

    for column_index in range(1, shipment_column_count + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 16
    for column_index in range(shipment_column_count + 1, total_columns + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 11


def _write_rate_card_sheet_inplace(
    worksheet,
    matrix: MatrixWorkbookData,
    *,
    changed_cells: set[tuple[int, int]] | None = None,
    updated_rows: set[int] | None = None,
) -> None:
    changed = changed_cells or set()
    highlighted_rows = updated_rows or set()
    total_columns = len(matrix.columns)

    first_clear_row = matrix.data_start_row + len(matrix.lanes)
    write_columns = len(matrix.columns)

    for lane_offset, lane in enumerate(matrix.lanes):
        row_index = matrix.data_start_row + lane_offset
        row_is_updated = row_index in highlighted_rows or lane.is_updated_row
        for column_index in range(1, write_columns + 1):
            value = lane.values[column_index - 1] if column_index - 1 < len(lane.values) else None
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if value == "":
                cell.value = None
            if row_is_updated:
                cell.fill = _updated_row_fill()
            if (row_index, column_index) in changed:
                cell.fill = _updated_fill()

    for row_index in range(first_clear_row, worksheet.max_row + 1):
        for column_index in range(1, write_columns + 1):
            worksheet.cell(row=row_index, column=column_index, value=None)

    lane_number_index = find_shipment_column_index(matrix.columns, "Lane #")
    if lane_number_index is not None:
        for lane_offset, lane in enumerate(matrix.lanes):
            row_index = matrix.data_start_row + lane_offset
            column_index = lane_number_index + 1
            lane.values[lane_number_index] = lane_offset + 1
            cell = worksheet.cell(row=row_index, column=column_index, value=lane_offset + 1)
            row_is_updated = row_index in highlighted_rows or lane.is_updated_row
            if row_is_updated:
                cell.fill = _updated_row_fill()
            if (row_index, column_index) in changed:
                cell.fill = _updated_fill()


def write_matrix_workbook(
    matrix: MatrixWorkbookData,
    output_path: Path,
    *,
    changed_cells: set[tuple[int, int]] | None = None,
    updated_rows: set[int] | None = None,
    source_workbook_path: Path | None = None,
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if source_workbook_path is not None:
        import shutil

        shutil.copy2(source_workbook_path, output_path)
        workbook = load_workbook(output_path)
        sheet_name = matrix.sheet_name
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found in {source_workbook_path.name}."
            )
        worksheet = workbook[sheet_name]
        _write_rate_card_sheet_inplace(
            worksheet,
            matrix,
            changed_cells=changed_cells,
            updated_rows=updated_rows,
        )
        workbook.save(output_path)
        return output_path

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = matrix.sheet_name[:31]
    _write_matrix_sheet(worksheet, matrix, changed_cells=changed_cells)

    for sheet_name, sheet_df in matrix.extra_sheets.items():
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]
        extra_sheet = workbook.create_sheet(title=sheet_name[:31])
        for column_index, header in enumerate(sheet_df.columns, start=1):
            extra_sheet.cell(row=1, column=column_index, value=header)
        for row_index, (_, row) in enumerate(sheet_df.iterrows(), start=2):
            for column_index, header in enumerate(sheet_df.columns, start=1):
                extra_sheet.cell(row=row_index, column=column_index, value=row[header])

    workbook.save(output_path)
    return output_path
