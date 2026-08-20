"""
Build a matrix-format workbook from a Caterpillar Air processing dataframe.

Reads AIR_RATES_ENRICHED from processing/ and writes a matrix workbook to output/.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from build_accessorial_costs import (
    ACCESSORIAL_COLUMNS,
    ACCESSORIAL_SHEET_NAME,
    DELIVERY_CHARGES_SHEET,
    NUMERIC_ACCESSORIAL_COLUMNS,
    AccessorialBuildResult,
    build_accessorial_costs_from_delivery_charges,
)
from number_utils import (
    EXCEL_NUMBER_FORMAT,
    EXCEL_TEXT_FORMAT,
    format_number_as_dot_text,
    normalize_dataframe_numbers,
    normalize_numeric_columns,
    normalize_numeric_value,
)
from export_origin_postal_code_zones import (
    OriginPostalCodeZoneExportResult,
    OriginPostalCodeZoneResolver,
    default_postal_zones_path,
    export_origin_postal_code_zones,
)
from project_paths import OUTPUT_DIR, PROCESSING_DIR, ensure_workspace_dirs

DEFAULT_SHEET_NAME = "AIR_RATES_ENRICHED"
EXCEL_SUFFIXES = {".xlsx", ".xls", ".xlsm"}

COST_NAME_ROW = 1
RATE_BY_ROW = 2
COLUMN_HEADER_ROW = 3
DATA_START_ROW = 4

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SUBHEADER_FONT = Font(bold=True, color="1F4E78")
BOLD_FONT = Font(bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="B4B4B4"),
    right=Side(style="thin", color="B4B4B4"),
    top=Side(style="thin", color="B4B4B4"),
    bottom=Side(style="thin", color="B4B4B4"),
)
PRICE_NUMBER_FORMAT = EXCEL_NUMBER_FORMAT

BOLD_SHIPMENT_HEADERS = {
    "Commodity",
    "Origin Port",
    "Origin Postal Code Zone",
    "Origin Postal Code",
    "Origin Country",
}

COUNTRY_PREFIX_PATTERN = re.compile(r"^[A-Za-z]{2}(.+)$")


@dataclass(frozen=True)
class ShipmentColumn:
    header: str
    source_column: str | None = None
    value_fn: str | None = None
    bold_header: bool = False


@dataclass(frozen=True)
class CostValueColumn:
    header: str
    source_column: str
    rate_unit: str = "Flat"


@dataclass
class CostBlock:
    cost_name: str
    currency_column: str
    rate_column: str
    min_column: str | None = None
    max_column: str | None = None
    value_columns: list[CostValueColumn] = field(default_factory=list)


@dataclass
class MatrixBuildResult:
    matrix_path: Path
    row_count: int
    shipment_column_count: int
    cost_block_count: int
    accessorial: AccessorialBuildResult | None = None
    postal_zones: OriginPostalCodeZoneExportResult | None = None


SHIPMENT_COLUMNS: tuple[ShipmentColumn, ...] = (
    ShipmentColumn("Lane #", value_fn="lane_number"),
    ShipmentColumn("KEY", source_column="KEY"),
    ShipmentColumn("Commodity", source_column="COMMODITY", bold_header=True),
    ShipmentColumn("ORIGIN_LOCATION_NAME", source_column="ORIGIN_LOCATION_NAME__C"),
    ShipmentColumn("Origin Port", value_fn="empty", bold_header=True),
    ShipmentColumn(
        "Origin Postal Code Zone",
        value_fn="origin_postal_code_zone",
        bold_header=True,
    ),
    ShipmentColumn(
        "Origin Postal Code",
        source_column="ORIGIN_LOCATION_NAME__C",
        value_fn="trim_country_prefix",
        bold_header=True,
    ),
    ShipmentColumn("Origin Country", source_column="ORIGIN_COUNTRY__C", bold_header=True),
    ShipmentColumn(
        "Destination Port",
        source_column="DESTINATION_LOCATION_NAME__C",
        value_fn="trim_country_prefix",
    ),
    ShipmentColumn("Destination country", source_column="DESTINATION_COUNTRY__C"),
    ShipmentColumn("Service level", source_column="SERVICE_LEVEL"),
    ShipmentColumn("Service", source_column="SERVICE"),
    ShipmentColumn("OPERATIONAL_FLOW", value_fn="empty"),
    ShipmentColumn("Valid from", source_column="RATE_EFFECTIVE_DATE__C", value_fn="date_dd_mm_yyyy"),
    ShipmentColumn("Valid to", source_column="RATE_EXPIRATION_DATE__C", value_fn="date_dd_mm_yyyy"),
)


def _cell_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def _is_empty_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    return _cell_text(value) == ""


def trim_country_prefix(value: object) -> str:
    text = _cell_text(value)
    if not text or "," in text:
        return ""
    match = COUNTRY_PREFIX_PATTERN.match(text)
    if not match:
        return ""
    return match.group(1)


def format_date_dd_mm_yyyy(value: object) -> str | None:
    if _is_empty_value(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return None
    return parsed.strftime("%d.%m.%Y")


def list_processing_files() -> list[Path]:
    return [
        path
        for path in sorted(PROCESSING_DIR.iterdir())
        if path.is_file()
        and path.suffix.lower() in EXCEL_SUFFIXES
        and not path.name.startswith("~$")
    ]


def prompt_file_selection(files: list[Path], title: str) -> Path:
    if not files:
        raise FileNotFoundError(f"No Excel files found in: {PROCESSING_DIR}")

    if len(files) == 1:
        print(f"\nUsing only file in processing/: {files[0].name}")
        return files[0]

    print(f"\n{title}")
    for index, path in enumerate(files, start=1):
        print(f"  {index}. {path.name}")

    while True:
        raw = input("Enter file number: ").strip()
        try:
            choice = int(raw) - 1
            if 0 <= choice < len(files):
                return files[choice]
        except ValueError:
            pass
        print(f"Please enter a number between 1 and {len(files)}.")


def _shipment_value(
    row: pd.Series,
    column: ShipmentColumn,
    *,
    lane_number: int,
    zone_resolver: OriginPostalCodeZoneResolver | None = None,
) -> object:
    if column.value_fn == "lane_number":
        return lane_number
    if column.value_fn == "empty":
        return None
    if column.value_fn == "trim_country_prefix":
        source = column.source_column or ""
        return trim_country_prefix(row.get(source))
    if column.value_fn == "origin_postal_code_zone":
        if zone_resolver is None:
            return None
        return zone_resolver.resolve(row.name)
    if column.value_fn == "date_dd_mm_yyyy":
        source = column.source_column or ""
        return format_date_dd_mm_yyyy(row.get(source))

    source = column.source_column
    if source is None:
        return None
    value = row.get(source)
    return None if _is_empty_value(value) else value


def discover_cost_blocks(columns: list[str]) -> list[CostBlock]:
    blocks: list[CostBlock] = []
    seen_prefixes: set[str] = set()

    for column in columns:
        if not column.endswith(" Currency"):
            continue

        prefix = column[: -len(" Currency")]
        if prefix in seen_prefixes:
            continue
        seen_prefixes.add(prefix)

        rate_column = prefix if prefix == "Base Rate" else f"{prefix} Rate"
        if rate_column not in columns:
            continue

        min_column = f"{prefix} Min" if f"{prefix} Min" in columns else None
        max_column = f"{prefix} Max" if f"{prefix} Max" in columns else None

        value_columns = [CostValueColumn("Flat", rate_column, "Flat")]
        if min_column:
            value_columns.insert(0, CostValueColumn("Min", min_column, "Min"))
        if max_column:
            value_columns.append(CostValueColumn("Max", max_column, "Max"))

        blocks.append(
            CostBlock(
                cost_name=rate_column,
                currency_column=column,
                rate_column=rate_column,
                min_column=min_column,
                max_column=max_column,
                value_columns=value_columns,
            )
        )

    return blocks


def _expand_block_columns(block: CostBlock) -> list[tuple[str, str, str, bool]]:
    expanded: list[tuple[str, str, str, bool]] = [
        ("Currency", block.currency_column, "Currency", True),
    ]
    for value_column in block.value_columns:
        expanded.append(
            (value_column.header, value_column.source_column, value_column.rate_unit, False)
        )
    return expanded


def _sanitize_cost_value(value: object) -> object:
    return normalize_numeric_value(value)


def _cost_block_has_data(block: CostBlock, rate_card: pd.DataFrame) -> bool:
    for _, row in rate_card.iterrows():
        for _, source_column, _, _ in _expand_block_columns(block):
            if source_column not in row.index:
                continue
            if not _is_empty_value(row[source_column]):
                return True
    return False


def _build_matrix_rows(
    rate_card: pd.DataFrame,
    shipment_columns: tuple[ShipmentColumn, ...],
    cost_blocks: list[CostBlock],
    *,
    zone_resolver: OriginPostalCodeZoneResolver,
) -> list[list[object]]:
    shipment_headers = [column.header for column in shipment_columns]
    expanded_cost_columns = [
        (block, column_def)
        for block in cost_blocks
        for column_def in _expand_block_columns(block)
    ]

    header_rows: list[list[object]] = []
    for header_index in range(3):
        row_values: list[object] = []
        if header_index == 2:
            row_values.extend(shipment_headers)
        else:
            row_values.extend([""] * len(shipment_headers))

        for block in cost_blocks:
            block_columns = _expand_block_columns(block)
            for header_label, _, rate_unit, is_currency in block_columns:
                if header_index == 0:
                    row_values.append(block.cost_name)
                elif header_index == 1:
                    row_values.append("")
                else:
                    row_values.append("Currency" if is_currency else header_label)

        header_rows.append(row_values)

    data_rows: list[list[object]] = []
    for lane_number, (_, row) in enumerate(rate_card.iterrows(), start=1):
        data_row: list[object] = []
        for column in shipment_columns:
            data_row.append(
                _shipment_value(
                    row,
                    column,
                    lane_number=lane_number,
                    zone_resolver=zone_resolver,
                )
            )

        for block, (_, source_column, _, is_currency) in [
            (block, column_def)
            for block in cost_blocks
            for column_def in _expand_block_columns(block)
        ]:
            if source_column not in row.index:
                data_row.append(None)
                continue
            value = row[source_column]
            if is_currency:
                data_row.append(None if _is_empty_value(value) else value)
            else:
                data_row.append(_sanitize_cost_value(value))

        data_rows.append(data_row)

    return header_rows + data_rows


def _merge_cost_name_cells(
    worksheet,
    shipment_column_count: int,
    cost_blocks: list[CostBlock],
) -> None:
    current_column = shipment_column_count + 1
    for block in cost_blocks:
        block_width = len(_expand_block_columns(block))
        if block_width > 1:
            worksheet.merge_cells(
                start_row=COST_NAME_ROW,
                start_column=current_column,
                end_row=COST_NAME_ROW,
                end_column=current_column + block_width - 1,
            )
        current_column += block_width


def _apply_worksheet_formatting(
    worksheet,
    shipment_columns: tuple[ShipmentColumn, ...],
    cost_blocks: list[CostBlock],
    total_rows: int,
) -> None:
    shipment_column_count = len(shipment_columns)
    cost_column_count = sum(len(_expand_block_columns(block)) for block in cost_blocks)
    total_columns = shipment_column_count + cost_column_count
    header_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left = Alignment(horizontal="left", vertical="center", wrap_text=False)

    for row_index in range(1, total_rows + 1):
        for column_index in range(1, total_columns + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.border = THIN_BORDER

            if row_index == COST_NAME_ROW and column_index > shipment_column_count:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = header_center
            elif row_index == RATE_BY_ROW:
                cell.alignment = center
            elif row_index == COLUMN_HEADER_ROW:
                cell.fill = SUBHEADER_FILL
                cell.font = SUBHEADER_FONT
                cell.alignment = center
                if column_index <= shipment_column_count:
                    shipment_column = shipment_columns[column_index - 1]
                    if shipment_column.bold_header:
                        cell.font = BOLD_FONT
            elif row_index >= DATA_START_ROW and column_index <= shipment_column_count:
                cell.alignment = left
            elif row_index >= DATA_START_ROW and column_index > shipment_column_count:
                cell.alignment = center
                cell.number_format = PRICE_NUMBER_FORMAT


def _write_accessorial_sheet(workbook, accessorial_df: pd.DataFrame) -> AccessorialBuildResult:
    if ACCESSORIAL_SHEET_NAME in workbook.sheetnames:
        del workbook[ACCESSORIAL_SHEET_NAME]

    accessorial_df = normalize_numeric_columns(accessorial_df, list(NUMERIC_ACCESSORIAL_COLUMNS))

    worksheet = workbook.create_sheet(title=ACCESSORIAL_SHEET_NAME[:31])
    for column_index, header in enumerate(ACCESSORIAL_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.fill = SUBHEADER_FILL
        cell.font = SUBHEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for row_index, (_, row) in enumerate(accessorial_df.iterrows(), start=2):
        for column_index, header in enumerate(ACCESSORIAL_COLUMNS, start=1):
            value = row[header]
            if header in NUMERIC_ACCESSORIAL_COLUMNS:
                value = format_number_as_dot_text(value)
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.border = THIN_BORDER
            if header in NUMERIC_ACCESSORIAL_COLUMNS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if value not in (None, ""):
                    cell.number_format = EXCEL_TEXT_FORMAT
            elif header == "Currency":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for column_index, header in enumerate(ACCESSORIAL_COLUMNS, start=1):
        width = 36 if header == "Applies if" else 14
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    return AccessorialBuildResult(row_count=len(accessorial_df))


def build_matrix_from_rate_card(
    rate_card: pd.DataFrame,
    output_path: Path,
    *,
    sheet_name: str = "Matrix",
    delivery_charges: pd.DataFrame | None = None,
) -> MatrixBuildResult:
    cost_blocks = discover_cost_blocks(list(rate_card.columns))
    cost_blocks = [block for block in cost_blocks if _cost_block_has_data(block, rate_card)]

    if not cost_blocks:
        raise ValueError("No cost blocks with data found in rate card.")

    zone_resolver = OriginPostalCodeZoneResolver(rate_card)
    matrix_rows = _build_matrix_rows(
        rate_card,
        SHIPMENT_COLUMNS,
        cost_blocks,
        zone_resolver=zone_resolver,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]

    for row_index, row_values in enumerate(matrix_rows, start=1):
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            if value == "":
                cell.value = None

    _merge_cost_name_cells(worksheet, len(SHIPMENT_COLUMNS), cost_blocks)
    _apply_worksheet_formatting(
        worksheet,
        SHIPMENT_COLUMNS,
        cost_blocks,
        len(matrix_rows),
    )

    for column_index in range(1, len(SHIPMENT_COLUMNS) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 16

    cost_offset = len(SHIPMENT_COLUMNS)
    for offset in range(sum(len(_expand_block_columns(block)) for block in cost_blocks)):
        letter = get_column_letter(cost_offset + offset + 1)
        worksheet.column_dimensions[letter].width = 11

    accessorial_result: AccessorialBuildResult | None = None
    if delivery_charges is not None and not delivery_charges.empty:
        accessorial_df = build_accessorial_costs_from_delivery_charges(delivery_charges)
        if not accessorial_df.empty:
            accessorial_result = _write_accessorial_sheet(workbook, accessorial_df)

    workbook.save(output_path)

    postal_zones_result = export_origin_postal_code_zones(
        rate_card,
        default_postal_zones_path(output_path),
        resolver=zone_resolver,
    )

    return MatrixBuildResult(
        matrix_path=output_path,
        row_count=len(rate_card),
        shipment_column_count=len(SHIPMENT_COLUMNS),
        cost_block_count=len(cost_blocks),
        accessorial=accessorial_result,
        postal_zones=postal_zones_result,
    )


def load_processing_sheet(
    processing_file: Path,
    sheet_name: str,
) -> pd.DataFrame | None:
    workbook = pd.ExcelFile(processing_file)
    if sheet_name not in workbook.sheet_names:
        return None
    df = pd.read_excel(processing_file, sheet_name=sheet_name)
    return normalize_dataframe_numbers(df)


def load_rate_card_from_processing(
    processing_file: Path,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> pd.DataFrame:
    workbook = pd.ExcelFile(processing_file)
    if sheet_name not in workbook.sheet_names:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in {processing_file.name}. "
            f"Available: {', '.join(workbook.sheet_names)}"
        )
    return pd.read_excel(processing_file, sheet_name=sheet_name)


def run_build_matrix(
    *,
    auto: bool = False,
    processing_file: Path | None = None,
    sheet_name: str = DEFAULT_SHEET_NAME,
    output_path: Path | None = None,
) -> MatrixBuildResult:
    ensure_workspace_dirs()

    processing_files = list_processing_files()
    selected_file = processing_file or (
        processing_files[0] if auto and processing_files else prompt_file_selection(
            processing_files,
            "Select processing file:",
        )
    )

    if not selected_file.exists():
        raise FileNotFoundError(f"Processing file not found: {selected_file}")

    print(f"\nLoading '{sheet_name}' from {selected_file.name}...")
    rate_card = load_rate_card_from_processing(selected_file, sheet_name=sheet_name)
    print(f"  {len(rate_card)} rows, {len(rate_card.columns)} columns")

    delivery_charges = load_processing_sheet(selected_file, DELIVERY_CHARGES_SHEET)
    if delivery_charges is not None:
        print(f"  DELIVERY_CHARGES: {len(delivery_charges)} rows")
    else:
        print("  DELIVERY_CHARGES: not found (skipping Accessorial cost tab)")

    final_output = output_path or (
        OUTPUT_DIR / f"{selected_file.stem}_matrix.xlsx"
    )

    print("\nBuilding matrix workbook...")
    result = build_matrix_from_rate_card(
        rate_card,
        final_output,
        delivery_charges=delivery_charges,
    )
    print(f"Saved matrix workbook to: {result.matrix_path}")
    print(
        f"  Rows: {result.row_count} | "
        f"Shipment columns: {result.shipment_column_count} | "
        f"Cost blocks: {result.cost_block_count}"
    )
    if result.accessorial is not None:
        print(
            f"  Accessorial tab '{result.accessorial.sheet_name}': "
            f"{result.accessorial.row_count} rows"
        )
    if result.postal_zones is not None:
        print(
            f"  Origin postal zones: {result.postal_zones.output_path} "
            f"({result.postal_zones.zone_count} zones)"
        )
    return result


def main() -> int:
    try:
        run_build_matrix()
        return 0
    except KeyboardInterrupt:
        print("\nOperation cancelled.")
        return 130
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
